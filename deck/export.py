# usage: python deck/export.py deck   (run from the repo root)
# render deck slide files -> PNG (headless chromium) -> PDF + picture pptx with notes
import json, re, sys, html, subprocess, glob, os
S = os.path.abspath(sys.argv[1]); D = f"{S}/project"; R = f"{S}/render"; os.makedirs(R, exist_ok=True)
# /_blob/<id> asset ids in the slides -> image files under S/img/
blobs = {"939e9d12fa84ebadfb1b47c68779fc23": "dash-live.png", "1c5bf139d362544fec913b712325d378": "calib-camera.png",
         "8624698d1c05014614ebc996a92f0931": "bike-lock.jpg", "85f0303797ef2f847365c1ac04128591": "catalytic.jpg",
         "39fb0b2e7b6874c21b876745f4901703": "fence-climb.jpg"}
deck = json.load(open(f"{D}/deck.json"))
links = "".join(f'<link rel="stylesheet" href="{f["href"]}">' for f in deck["faces"].values())
arrow = ('<svg width="48" height="24" viewBox="0 0 48 24" style="flex:none"><line x1="0" y1="12" x2="38" y2="12" '
         'stroke="#FF6A1A" stroke-width="3"/><path d="M36 4 L48 12 L36 20 Z" fill="#FF6A1A"/></svg>')
css = ("*{margin:0;box-sizing:border-box} html,body{width:1920px;height:1080px;overflow:hidden;background:#0B0C0F}"
       "section{position:relative;width:1920px;height:1080px;overflow:hidden} aside{display:none}"
       "table{border-collapse:collapse;width:100%} th,td{padding:0.35em 0.6em;border-bottom:1px solid #2A2D34;text-align:left;font-weight:600}"
       "td{font-weight:400} h1,h2,h3{font-weight:600}")
chrome = glob.glob(os.path.expanduser("~/.cache/ms-playwright/chromium-1247/*/chrome"))[0]
for sid in deck["order"]:
    h = open(f"{D}/slides/{sid}.html").read()
    h = re.sub(r"<x-connector[^>]*></x-connector>", arrow, h)
    for b, f in blobs.items():
        h = h.replace(f"/_blob/{b}", "file://" + os.path.abspath(f"{S}/img/{f}"))
    open(f"{R}/{sid}.html", "w").write(f"<!doctype html><html><head><meta charset='utf-8'>{links}<style>{css}</style></head><body>{h}</body></html>")
    subprocess.run([chrome, "--headless=new", "--no-sandbox", "--disable-gpu", "--hide-scrollbars", "--allow-file-access-from-files",
                    f"--user-data-dir={S}/chrome-prof", "--window-size=1920,1080", "--virtual-time-budget=6000",
                    f"--screenshot={R}/{sid}.png", f"file://{R}/{sid}.html"], capture_output=True, timeout=60, check=True)
from PIL import Image
from pptx import Presentation
from pptx.util import Emu
ims = [Image.open(f"{R}/{s}.png").convert("RGB") for s in deck["order"]]
ims[0].save(f"{S}/Sauron-Hackathon-Deck.pdf", save_all=True, append_images=ims[1:], resolution=144)
p = Presentation(); p.slide_width = Emu(12192000); p.slide_height = Emu(6858000)
for s in deck["order"]:
    sl = p.slides.add_slide(p.slide_layouts[6])
    sl.shapes.add_picture(f"{R}/{s}.png", 0, 0, p.slide_width, p.slide_height)
    m = re.search(r"<aside>(.*?)</aside>", open(f"{D}/slides/{s}.html").read(), re.S)
    if m: sl.notes_slide.notes_text_frame.text = html.unescape(m.group(1))
p.save(f"{S}/Sauron-Hackathon-Deck.pptx")
print("slides", len(deck["order"]))
